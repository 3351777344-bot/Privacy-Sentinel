"""Build the reviewed Word acceptance report from the repository Markdown."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root = Path(__file__).resolve().parents[1]
source = root / 'platform/docs/验收报告.md'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2)
for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = document.styles[style_name]
    style.font.name = 'Microsoft YaHei'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
for node in list(document.styles.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
normal = document.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.2
normal.paragraph_format.space_after = Pt(7)
document.styles['Title'].font.size = Pt(22)
document.styles['Heading 1'].font.size = Pt(16)
document.styles['Heading 1'].paragraph_format.space_after = Pt(12)
first_section = True
lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith('# '):
        document.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        if not first_section:
            document.add_page_break()
        document.add_heading(line[3:], 1)
        first_section = False
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [x.strip() for x in lines[i].strip('|').split('|')]
            if not all(set(x) <= set('-: ') for x in cells):
                rows.append(cells)
            i += 1
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for column, width in zip(table.columns, [3.5, 6.0, 7.5]):
            column.width = Cm(width)
        for n, row in enumerate(rows):
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = value
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for run in cells[index].paragraphs[0].runs:
                    run.font.size = Pt(9)
                    run.bold = n == 0
                properties = cells[index]._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'DCE6F1' if n == 0 else ('F5F7FA' if n % 2 == 0 else 'FFFFFF'))
                properties.append(shade)
                borders = OxmlElement('w:tcBorders')
                for edge in ['top', 'left', 'bottom', 'right']:
                    node = OxmlElement('w:' + edge)
                    node.set(qn('w:val'), 'single')
                    node.set(qn('w:sz'), '4')
                    node.set(qn('w:color'), 'D9D9D9')
                    borders.append(node)
                properties.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge in ['top', 'left', 'bottom', 'right']:
                    node = OxmlElement('w:' + edge)
                    node.set(qn('w:w'), '90')
                    node.set(qn('w:type'), 'dxa')
                    margins.append(node)
                properties.append(margins)
            if n == 0:
                repeat = OxmlElement('w:tblHeader')
                table.rows[n]._tr.get_or_add_trPr().append(repeat)
        document.add_paragraph()
        continue
    elif line.strip():
        document.add_paragraph(line)
    i += 1
document.core_properties.title = 'GuardianHub 检测优化与验收报告'
document.core_properties.subject = '工程验证与比赛发布条件'
document.core_properties.author = 'GuardianHub 工程验收'
document.save(root / 'platform/docs/GuardianHub 验收报告.docx')
print('Word acceptance report generated.')
